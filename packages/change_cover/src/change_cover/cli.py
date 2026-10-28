from __future__ import annotations

import argparse
from pathlib import Path
from typing import Sequence

from pptx import Presentation
from pptx.parts.image import Image


DEFAULT_AUTHOR = "Victor Chavarrias\nWillem Ottevanger"


def _expand_line_breaks(value: str) -> str:
    return value.replace(r"\n", "\n")


def update_presentation(
    input_path: Path,
    output_path: Path,
    *,
    title: str = "",
    author: str = DEFAULT_AUTHOR,
    old_author: str = "Victor",
    image_path: Path | None = None,
) -> None:
    """Update the first slide of one presentation and save it to ``output_path``."""
    presentation = Presentation(input_path)
    first_slide = presentation.slides[0]
    new_title = _expand_line_breaks(title)
    new_author = _expand_line_breaks(author)

    if new_title and first_slide.shapes.title:
        first_slide.shapes.title.text = new_title

    replacement_blob = (
        Image.from_file(str(image_path))._blob if image_path is not None else None
    )
    for shape in first_slide.shapes:
        if (
            old_author
            and new_author
            and shape.has_text_frame
            and old_author in shape.text
        ):
            alignment = shape.text_frame.paragraphs[0].alignment
            shape.text = new_author
            for paragraph in shape.text_frame.paragraphs:
                paragraph.alignment = alignment

        if replacement_blob is not None and not shape.has_text_frame:
            image_rid = getattr(shape._element, "blip_rId", None)
            if image_rid:
                image_part = shape.part.related_part(image_rid)
                image_part.blob = replacement_blob

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def update_directory(
    input_dir: Path,
    output_dir: Path,
    *,
    title: str = "",
    author: str = DEFAULT_AUTHOR,
    old_author: str = "Victor",
    image_path: Path | None = None,
) -> int:
    """Update all PowerPoint files in ``input_dir`` and return the file count."""
    input_dir = input_dir.resolve()
    output_dir = output_dir.resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input directory does not exist: {input_dir}")
    if image_path is not None:
        image_path = image_path.resolve()
        if not image_path.is_file():
            raise FileNotFoundError(f"Image file does not exist: {image_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    processed_files = 0
    for input_path in sorted(input_dir.iterdir()):
        if (
            not input_path.is_file()
            or input_path.name.startswith("~")
            or input_path.suffix.lower() != ".pptx"
        ):
            continue
        update_presentation(
            input_path,
            output_dir / input_path.name,
            title=title,
            author=author,
            old_author=old_author,
            image_path=image_path,
        )
        processed_files += 1
    return processed_files


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    """Parse command-line arguments for the ``change-cover`` command."""
    parser = argparse.ArgumentParser(
        description="Update the title, author, and cover image in PowerPoint files."
    )
    parser.add_argument(
        "--title",
        default="",
        help="New title; leave empty to keep the existing title.",
    )
    parser.add_argument(
        "--author",
        default=DEFAULT_AUTHOR,
        help=r"New author; leave empty to keep the existing author. Use \n for line breaks.",
    )
    parser.add_argument(
        "--old-author",
        default="Victor",
        help="Text used to identify the existing author.",
    )
    parser.add_argument(
        "--image",
        type=Path,
        default=None,
        help="Path to the replacement image; leave empty to keep the existing image.",
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=Path.cwd() / "day4",
        help="Directory containing the PowerPoint files (default: ./day4).",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path.cwd() / "day4_modified",
        help="Directory where modified files are written (default: ./day4_modified).",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    processed_files = update_directory(
        args.input_dir,
        args.output_dir,
        title=args.title,
        author=args.author,
        old_author=args.old_author,
        image_path=args.image,
    )
    print(f"Updated {processed_files} PowerPoint file(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())