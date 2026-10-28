from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from change_cover import cli
from change_cover import update_directory


def _create_presentation(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Original title"
    text_box = slide.shapes.add_textbox(
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(1),
    )
    text_box.text = "Victor is the original author"
    presentation.save(path)


def test_update_directory_changes_title_and_author(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    _create_presentation(input_dir / "deck.pptx")
    (input_dir / "~temporary.pptx").write_bytes(b"not a presentation")
    (input_dir / "notes.txt").write_text("ignored", encoding="utf-8")

    processed = update_directory(
        input_dir,
        output_dir,
        title=r"Updated\nTitle",
        author=r"New Author",
    )

    assert processed == 1
    result = Presentation(output_dir / "deck.pptx")
    first_slide = result.slides[0]
    assert first_slide.shapes.title.text == "Updated\nTitle"
    assert any(shape.text == "New Author" for shape in first_slide.shapes if shape.has_text_frame)
    assert not (output_dir / "~temporary.pptx").exists()


def test_update_directory_requires_input_directory(tmp_path: Path) -> None:
    missing = tmp_path / "missing"

    try:
        update_directory(missing, tmp_path / "output")
    except FileNotFoundError as error:
        assert str(missing) in str(error)
    else:
        raise AssertionError("Expected FileNotFoundError")


def test_cli_delegates_to_python_api(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: dict[str, object] = {}

    def fake_update_directory(
        input_dir: Path,
        output_dir: Path,
        **kwargs: object,
    ) -> int:
        calls.update(
            input_dir=input_dir,
            output_dir=output_dir,
            **kwargs,
        )
        return 2

    monkeypatch.setattr(cli, "update_directory", fake_update_directory)

    result = cli.main(
        [
            "--input-dir",
            str(tmp_path / "input"),
            "--output-dir",
            str(tmp_path / "output"),
        ]
    )

    assert result == 0
    assert calls["input_dir"] == tmp_path / "input"
    assert calls["output_dir"] == tmp_path / "output"
